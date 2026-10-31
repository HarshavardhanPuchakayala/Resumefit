from io import BytesIO
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    HRFlowable,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


def _dates(start: str, end: str) -> str:
    return " – ".join(part for part in (start, end) if part)


def build_blocks(contact: dict, content: dict) -> list[tuple]:
    blocks: list[tuple] = []

    if contact.get("full_name"):
        blocks.append(("name", contact["full_name"]))
    if content.get("headline"):
        blocks.append(("headline", content["headline"]))

    line1 = " | ".join(
        part for part in (contact.get("location"), contact.get("email"), contact.get("phone")) if part
    )
    line2 = " | ".join(
        part for part in (contact.get("linkedin"), contact.get("github"), contact.get("portfolio")) if part
    )
    if line1:
        blocks.append(("line", line1))
    if line2:
        blocks.append(("line", line2))

    if content.get("summary"):
        blocks.append(("heading", "Professional Summary"))
        blocks.append(("text", content["summary"]))

    groups = [g for g in content.get("skills", []) if g.get("items")]
    if groups:
        blocks.append(("heading", "Technical Skills"))
        for group in groups:
            label = f"{group['category']}: " if group.get("category") else ""
            blocks.append(("text", label + ", ".join(group["items"])))

    if content.get("experience"):
        blocks.append(("heading", "Professional Experience"))
        for job in content["experience"]:
            title = " — ".join(part for part in (job.get("role"), job.get("company")) if part)
            if job.get("location"):
                title += f", {job['location']}"
            blocks.append(("row", title, _dates(job.get("start_date", ""), job.get("end_date", ""))))
            for bullet in job.get("bullets", []):
                blocks.append(("bullet", bullet))

    if content.get("projects"):
        blocks.append(("heading", "Projects"))
        for project in content["projects"]:
            blocks.append(
                ("row", project.get("name", ""), _dates(project.get("start_date", ""), project.get("end_date", "")))
            )
            if project.get("tech_stack"):
                blocks.append(("text", "Tech Stack: " + ", ".join(project["tech_stack"])))
            for bullet in project.get("bullets", []):
                blocks.append(("bullet", bullet))

    if content.get("education"):
        blocks.append(("heading", "Education"))
        for item in content["education"]:
            blocks.append(("row", item.get("degree", ""), _dates(item.get("start_year", ""), item.get("end_year", ""))))
            place = ", ".join(part for part in (item.get("institution"), item.get("location")) if part)
            if place:
                blocks.append(("text", place))

    return blocks


def _bottom_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "4"), ("w:space", "1"), ("w:color", "808080")):
        bottom.set(qn(key), value)
    borders.append(bottom)
    p_pr.append(borders)


def build_docx(blocks: list[tuple]) -> bytes:
    doc = Document()
    doc.core_properties.author = ""

    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(18)
    section.right_margin = Mm(18)
    section.top_margin = Mm(15)
    section.bottom_margin = Mm(15)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    usable = Mm(174)

    for block in blocks:
        kind = block[0]

        if kind == "name":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(block[1].upper())
            run.bold = True
            run.font.size = Pt(18)
        elif kind == "headline":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(block[1]).font.size = Pt(12)
        elif kind == "line":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(block[1]).font.size = Pt(9.5)
        elif kind == "heading":
            p = doc.add_paragraph()
            run = p.add_run(block[1].upper())
            run.bold = True
            run.font.size = Pt(11)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            _bottom_border(p)
        elif kind == "row":
            p = doc.add_paragraph()
            p.paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
            p.add_run(block[1]).bold = True
            p.add_run("\t" + block[2])
        elif kind == "bullet":
            doc.add_paragraph(block[1], style="List Bullet")
        else:
            doc.add_paragraph(block[1])

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_pdf(blocks: list[tuple]) -> bytes:
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title="Resume",
        author="",
    )
    usable = A4[0] - 36 * mm

    name_style = ParagraphStyle("name", fontName="Helvetica-Bold", fontSize=18, leading=22, alignment=TA_CENTER)
    headline_style = ParagraphStyle("headline", fontName="Helvetica", fontSize=11, leading=14, alignment=TA_CENTER)
    line_style = ParagraphStyle("line", fontName="Helvetica", fontSize=9, leading=12, alignment=TA_CENTER)
    section_style = ParagraphStyle("section", fontName="Helvetica-Bold", fontSize=10.5, leading=13, spaceBefore=8)
    text_style = ParagraphStyle("text", fontName="Helvetica", fontSize=9.5, leading=12.5)
    bold_style = ParagraphStyle("bold", parent=text_style, fontName="Helvetica-Bold")
    right_style = ParagraphStyle("right", parent=text_style, alignment=TA_RIGHT)
    bullet_style = ParagraphStyle("bullet", parent=text_style, leftIndent=14, bulletIndent=4)

    row_style = TableStyle(
        [
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]
    )

    story = []
    for block in blocks:
        kind = block[0]

        if kind == "name":
            story.append(Paragraph(escape(block[1].upper()), name_style))
        elif kind == "headline":
            story.append(Paragraph(escape(block[1]), headline_style))
        elif kind == "line":
            story.append(Paragraph(escape(block[1]), line_style))
        elif kind == "heading":
            story.append(Paragraph(escape(block[1].upper()), section_style))
            story.append(HRFlowable(width="100%", thickness=0.6, color=colors.grey, spaceBefore=1, spaceAfter=3))
        elif kind == "row":
            table = Table(
                [[Paragraph(escape(block[1]), bold_style), Paragraph(escape(block[2]), right_style)]],
                colWidths=[usable * 0.72, usable * 0.28],
            )
            table.setStyle(row_style)
            story.append(table)
        elif kind == "bullet":
            story.append(Paragraph(escape(block[1]), bullet_style, bulletText="•"))
        else:
            story.append(Paragraph(escape(block[1]), text_style))

    if not story:
        story.append(Spacer(1, 1))

    doc.build(story)
    return buffer.getvalue()